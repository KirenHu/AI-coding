from pathlib import Path

from docx import Document

from worktwin.collector import Collector
from worktwin.db import Database
from worktwin.search import search
from tests.support import run_ai


def setup_db(tmp_path, source):
    db = Database(tmp_path / "worktwin.sqlite")
    with db.connect() as con:
        con.execute("INSERT INTO sources(name,kind,root,allow_ai) VALUES(?,?,?,1)",("工作材料","folder",str(source.resolve())))
    return db


def test_incremental_scan_search_and_delete(tmp_path: Path):
    folder = tmp_path / "authorized"
    folder.mkdir()
    md = folder / "decisions.md"
    md.write_text("# 审批配置\n\n经过权衡，我们最终决定把审批能力做成流程中的一个节点，避免独立审批流增加交互复杂度。\n\n我更倾向先给出结论，然后再说明必要的依据。",encoding="utf-8")
    (folder / "README.py").write_text("# coding: utf-8\nprint('hello')",encoding="utf-8")
    (folder / ".env").write_text("DATABASE_PASSWORD=secret")
    (folder / "secret.key").write_text("private key")
    outside = tmp_path / "outside.md"
    outside.write_text("不在授权目录内的文件")
    (folder / "outside-link.md").symlink_to(outside)
    db = setup_db(tmp_path, folder)
    collector = Collector(db)
    first = collector.scan_all()
    assert first["new"] == 2
    assert first["errors"] == 0
    assert collector.scan_all()["new"] == 0
    run_ai(db)
    with db.connect() as con:
        docs = con.execute("SELECT path,content FROM documents").fetchall()
        assert len(docs) == 2
        assert all("secret" not in r["content"] for r in docs)
        assert len(search(con, "审批")["documents"]) >= 1  # two-character Chinese LIKE fallback
        assert len(search(con, "最终决定")["documents"]) >= 1
        kinds = {r["kind"] for r in con.execute("SELECT kind FROM knowledge")}
        assert "decision" in kinds
        assert "preference" not in kinds  # Arbitrary documents cannot assert employee identity
        evidence = con.execute("SELECT COUNT(*) FROM knowledge_evidence").fetchone()[0]
        assert evidence >= 1
    md.write_text("# 审批配置\n\n现在最终决定按新的业务流程优化审批节点配置，取消原来的动态分支。",encoding="utf-8")
    assert collector.scan_all()["updated"] == 1
    with db.connect() as con:
        old = con.execute("SELECT COUNT(*) FROM knowledge WHERE needs_review=1").fetchone()[0]
        assert old >= 1
    md.unlink()
    assert collector.scan_all()["deleted"] == 1
    with db.connect() as con:
        assert not search(con,"动态分支")["documents"]
        assert con.execute("SELECT count(*) FROM documents").fetchone()[0] == 1


def test_docx_and_subfolders(tmp_path: Path):
    root = tmp_path / "root"
    folder = root / "manuals"
    folder.mkdir(parents=True)
    doc = Document()
    doc.add_heading("配置说明", level=1)
    doc.add_paragraph("我们最终选择人工复核，以保证产品安全。")
    doc.save(folder / "说明.docx")
    db = setup_db(tmp_path,root)
    result = Collector(db).scan_all()
    assert result["new"] == 1
    with db.connect() as con:
        assert len(search(con,"人工复核")["documents"]) == 1


def test_source_removal_purges_derived_data(tmp_path: Path):
    root = tmp_path / "w"
    root.mkdir()
    (root / "d.md").write_text("最终决定把所有工作流程记录为任务节点，并关联历史数据。",encoding="utf-8")
    db=setup_db(tmp_path,root)
    Collector(db).scan_all()
    run_ai(db)
    with db.connect() as con:
        assert con.execute("SELECT count(*) FROM knowledge").fetchone()[0] > 0
        con.execute("DELETE FROM sources")
        con.execute("DELETE FROM knowledge WHERE source_bound=1 AND NOT EXISTS (SELECT 1 FROM knowledge_evidence WHERE knowledge_id=knowledge.id)")
        assert con.execute("SELECT count(*) FROM documents").fetchone()[0] == 0
        assert con.execute("SELECT count(*) FROM chunks").fetchone()[0] == 0
        assert con.execute("SELECT count(*) FROM knowledge").fetchone()[0] == 0
