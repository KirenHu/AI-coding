"""Open-format and document text extraction using existing OSS libraries."""

from __future__ import annotations

import csv
import html
import json
import re
from html.parser import HTMLParser
from pathlib import Path

from .codex import parse_codex_session
from .claude import parse_claude_session
from .config import MAX_TEXT_CHARS

TEXT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".rst", ".log", ".csv", ".tsv",
    ".json", ".jsonl", ".yaml", ".yml", ".toml", ".html", ".htm",
    ".py", ".js", ".jsx", ".ts", ".tsx", ".vue", ".go", ".rs",
    ".java", ".sql", ".sh", ".css", ".xml", ".ipynb",
}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | {".docx", ".pdf"}
EXCLUDED_FILES = {".env", ".env.local", ".env.production", "id_rsa", "id_ed25519", "credentials.json", "secrets.json"}
EXCLUDED_FOLDERS = {".git", ".svn", ".hg", ".venv", "venv", "node_modules", "__pycache__", "dist", "build", ".next", ".idea", ".vscode", ".cache"}


class _VisibleHTML(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.out: list[str] = []
        self.hidden = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in ("script", "style", "noscript"):
            self.hidden += 1
        elif tag in ("p", "div", "li", "h1", "h2", "h3", "br"):
            self.out.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in ("script", "style", "noscript"):
            self.hidden = max(0, self.hidden - 1)

    def handle_data(self, text: str) -> None:
        if not self.hidden:
            self.out.append(text)


def parse_file(path: Path, *, codex: bool = False, claude: bool = False) -> tuple[str, str]:
    """Parse only the file explicitly supplied from an authorized source."""
    if codex:
        return parse_codex_session(path)
    if claude:
        return parse_claude_session(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        text = "\n\n".join((p.extract_text() or "") for p in reader.pages)
    elif suffix == ".docx":
        from docx import Document
        doc = Document(str(path))
        text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        for table in doc.tables:
            for row in table.rows:
                text += "\n" + " | ".join(cell.text.strip() for cell in row.cells)
    elif suffix == ".ipynb":
        obj = json.loads(path.read_text(encoding="utf-8", errors="replace"))
        cells = obj.get("cells", []) if isinstance(obj, dict) else []
        text = "\n\n".join("".join(c.get("source", [])) for c in cells if isinstance(c, dict) and c.get("cell_type") in ("markdown", "code"))
    elif suffix in (".html", ".htm"):
        parser = _VisibleHTML()
        parser.feed(path.read_text(encoding="utf-8", errors="replace"))
        text = html.unescape("".join(parser.out))
    elif suffix in (".csv", ".tsv"):
        separator = "\t" if suffix == ".tsv" else ","
        with path.open(encoding="utf-8-sig", errors="replace", newline="") as f:
            text = "\n".join(" | ".join(row) for _, row in zip(range(8000), csv.reader(f, delimiter=separator)))
    else:
        text = path.read_text(encoding="utf-8-sig", errors="replace")
    text = text.replace("\x00", "").strip()
    text = re.sub(r"\n{4,}", "\n\n\n", text)
    title_match = re.search(r"^#{1,2}\s+(.+)$", text, flags=re.MULTILINE)
    title = title_match.group(1).strip()[:100] if title_match else path.stem
    return title, text[:MAX_TEXT_CHARS]


def split_chunks(text: str, target_size: int = 1100, overlap: int = 120) -> list[tuple[int, int, str]]:
    """Create source-addressable chunks, preferring natural paragraph boundaries."""
    if not text.strip():
        return []
    result: list[tuple[int, int, str]] = []
    start = 0
    n = len(text)
    while start < n:
        limit = min(n, start + target_size)
        if limit < n:
            split = text.rfind("\n\n", start + target_size // 2, limit)
            if split < 0:
                split = text.rfind("\n", start + target_size // 2, limit)
            end = split + 2 if split >= 0 else limit
        else:
            end = n
        excerpt = text[start:end].strip()
        if excerpt:
            result.append((start, end, excerpt))
        if end == n:
            break
        start = max(start + 1, end - overlap)
    return result
